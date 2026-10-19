"""Excel and PDF files from a table on screen.

Every list in the app (payables, receivables, cash flow, registry, reports)
already holds the rows the user filtered. Rather than rebuilding each query
here, the page sends those rows and this module only lays them out — so the
file is exactly what was on screen.
"""

import io
import unicodedata
from datetime import date, datetime
from typing import Any, Iterable, List, Literal, Optional

from pydantic import BaseModel, Field

ColumnType = Literal["text", "money", "number", "date", "percent"]

MAX_ROWS = 20_000


class ExportColumn(BaseModel):
    key: str
    label: str
    type: ColumnType = "text"


class ExportRequest(BaseModel):
    title: str = Field(..., max_length=120)
    subtitle: Optional[str] = Field(None, max_length=300)
    columns: List[ExportColumn] = Field(..., min_length=1, max_length=30)
    rows: List[dict] = Field(default_factory=list, max_length=MAX_ROWS)
    # Linha final com totais, chaves iguais às das colunas.
    totals: Optional[dict] = None
    format: Literal["xlsx", "pdf"] = "xlsx"


def _as_number(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _as_date(value: Any) -> Optional[date]:
    if not value:
        return None
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


def _fmt_money(value: float) -> str:
    """1234.5 → '1 234,50' (formato português, sem símbolo)."""
    sign = "-" if value < 0 else ""
    whole, cents = f"{abs(value):.2f}".split(".")
    groups = []
    while whole:
        groups.insert(0, whole[-3:])
        whole = whole[:-3]
    return f"{sign}{' '.join(groups)},{cents}"


def _display(value: Any, col: ExportColumn) -> str:
    """Text for the PDF, where nothing is typed."""
    if col.type in ("money", "number", "percent"):
        num = _as_number(value)
        if num is None:
            return "" if value in (None, "") else str(value)
        if col.type == "money":
            return _fmt_money(num)
        if col.type == "percent":
            return f"{num:g}%".replace(".", ",")
        return f"{num:g}".replace(".", ",")
    if col.type == "date":
        d = _as_date(value)
        return d.strftime("%d/%m/%Y") if d else ("" if value is None else str(value))
    return "" if value is None else str(value)


def safe_filename(title: str, ext: str) -> str:
    """ASCII only: the name travels in an HTTP header ("Relatório" → "relatorio")."""
    plain = unicodedata.normalize("NFKD", title).encode("ascii", "ignore").decode("ascii")
    keep = "".join(c if c.isalnum() or c in " -_" else " " for c in plain)
    slug = "-".join(keep.lower().split()) or "exportacao"
    return f"{slug}-{date.today().isoformat()}.{ext}"


# ---------------------------------------------------------------- Excel

def to_xlsx(req: ExportRequest, company_name: str = "") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = req.title[:31] or "Dados"

    ncols = len(req.columns)
    ws.cell(row=1, column=1, value=req.title).font = Font(bold=True, size=14)
    info = " · ".join(p for p in (company_name, req.subtitle) if p)
    ws.cell(row=2, column=1, value=info or None).font = Font(color="666666", size=10)
    ws.cell(row=3, column=1, value=f"Exportado em {datetime.now().strftime('%d/%m/%Y %H:%M')}").font = Font(color="999999", size=9)

    header_row = 5
    head_fill = PatternFill("solid", fgColor="111111")
    thin = Side(style="thin", color="DDDDDD")
    for idx, col in enumerate(req.columns, start=1):
        cell = ws.cell(row=header_row, column=idx, value=col.label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(
            horizontal="right" if col.type in ("money", "number", "percent") else "left",
            vertical="center",
        )

    def write_row(r: int, data: dict, bold: bool = False):
        for c_idx, col in enumerate(req.columns, start=1):
            raw = data.get(col.key)
            if col.type in ("money", "number"):
                value = _as_number(raw)
                if value is None and raw not in (None, ""):
                    value = str(raw)
            elif col.type == "percent":
                num = _as_number(raw)
                value = num / 100 if num is not None else None
            elif col.type == "date":
                value = _as_date(raw) or (None if raw in (None, "") else str(raw))
            else:
                value = None if raw is None else str(raw)
            cell = ws.cell(row=r, column=c_idx, value=value)
            cell.border = Border(bottom=thin)
            if bold:
                cell.font = Font(bold=True)
            if col.type == "money":
                cell.number_format = '#,##0.00 "€";[Red]-#,##0.00 "€"'
            elif col.type == "percent":
                cell.number_format = "0.##%"
            elif col.type == "date":
                cell.number_format = "DD/MM/YYYY"

    r = header_row + 1
    for row in req.rows:
        write_row(r, row)
        r += 1
    if req.totals:
        write_row(r, req.totals, bold=True)

    # Larguras pela maior célula (limitadas, para descrições longas).
    for idx, col in enumerate(req.columns, start=1):
        longest = max(
            [len(col.label)] + [len(_display(row.get(col.key), col)) for row in req.rows[:500]]
        )
        ws.column_dimensions[get_column_letter(idx)].width = min(max(10, longest + 2), 60)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if req.rows:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncols)}{header_row + len(req.rows)}"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------- PDF

def _latin1(text: str) -> str:
    """As fontes base do PDF só têm latin-1; o euro passa a 'EUR'."""
    return text.replace("€", "EUR").replace("—", "-").replace("–", "-").replace("…", "...") \
        .replace("“", '"').replace("”", '"').replace("’", "'").replace("·", "-") \
        .encode("latin-1", "replace").decode("latin-1")


def _weights(req: ExportRequest) -> List[float]:
    """Column widths: text columns grow with their content, numbers stay narrow."""
    out = []
    for col in req.columns:
        sample: Iterable[str] = (_display(r.get(col.key), col) for r in req.rows[:300])
        longest = max([len(col.label)] + [len(s) for s in sample])
        if col.type == "text":
            out.append(min(max(longest, 8), 45))
        else:
            out.append(min(max(longest, 8), 16))
    return out


def to_pdf(req: ExportRequest, company_name: str = "") -> bytes:
    from fpdf import FPDF

    weights = _weights(req)
    landscape = sum(weights) > 110 or len(req.columns) > 6
    pdf = FPDF(orientation="L" if landscape else "P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.set_margins(12, 12, 12)
    pdf.alias_nb_pages()

    footer_text = _latin1(f"{company_name} - {req.title}" if company_name else req.title)

    def footer():
        pdf.set_y(-10)
        pdf.set_font("Helvetica", size=7)
        pdf.set_text_color(150, 150, 150)
        pdf.cell(0, 4, footer_text, align="L")
        pdf.set_x(pdf.l_margin)
        pdf.cell(0, 4, f"Página {pdf.page_no()}/{{nb}}".encode("latin-1", "replace").decode("latin-1"), align="R")

    pdf.footer = footer  # type: ignore[method-assign]
    pdf.add_page()

    width = pdf.w - pdf.l_margin - pdf.r_margin
    total_w = sum(weights)
    widths = [width * w / total_w for w in weights]
    line_h = 5.2

    pdf.set_font("Helvetica", "B", 14)
    pdf.set_text_color(17, 17, 17)
    pdf.cell(0, 7, _latin1(req.title), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=9)
    pdf.set_text_color(110, 110, 110)
    info = " - ".join(p for p in (company_name, req.subtitle) if p)
    if info:
        pdf.cell(0, 5, _latin1(info), new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 5, f"Exportado em {datetime.now().strftime('%d/%m/%Y %H:%M')} - valores em EUR",
             new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    def align(col: ExportColumn) -> str:
        return "R" if col.type in ("money", "number", "percent") else "L"

    def header():
        pdf.set_font("Helvetica", "B", 8)
        pdf.set_fill_color(17, 17, 17)
        pdf.set_text_color(255, 255, 255)
        for col, w in zip(req.columns, widths):
            pdf.cell(w, 6.5, _latin1(col.label), border=0, align=align(col), fill=True)
        pdf.ln()
        pdf.set_text_color(30, 30, 30)

    def fit(text: str, w: float) -> str:
        if pdf.get_string_width(text) <= w - 1.5:
            return text
        while text and pdf.get_string_width(text + "...") > w - 1.5:
            text = text[:-1]
        return text + "..."

    def row_out(data: dict, shade: bool, bold: bool = False):
        if pdf.get_y() + line_h > pdf.page_break_trigger:
            pdf.add_page()
            header()
        pdf.set_font("Helvetica", "B" if bold else "", 8)
        pdf.set_fill_color(245, 245, 245)
        for col, w in zip(req.columns, widths):
            text = fit(_latin1(_display(data.get(col.key), col)), w)
            pdf.cell(w, line_h, text, border="B" if bold else 0, align=align(col), fill=shade)
        pdf.ln()

    header()
    if not req.rows:
        pdf.set_font("Helvetica", "I", 9)
        pdf.cell(0, 8, "Sem dados para os filtros escolhidos.", new_x="LMARGIN", new_y="NEXT")
    for i, row in enumerate(req.rows):
        row_out(row, shade=i % 2 == 1)
    if req.totals:
        row_out(req.totals, shade=False, bold=True)

    return bytes(pdf.output())
