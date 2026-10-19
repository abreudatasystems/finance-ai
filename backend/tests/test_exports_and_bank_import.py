"""Exportar o que está no ecrã e importar o extrato como o banco o dá.

Os bancos portugueses entregam o extrato em Excel, com umas linhas sobre a
conta (titular, IBAN, período) antes dos títulos das colunas e um "Saldo
final" no fim. O leitor tinha de aceitar isso sem pedir a ninguém que
limpasse o ficheiro à mão.
"""

import io
from datetime import datetime

from openpyxl import Workbook, load_workbook

from app.services.bank_parser import parse_csv, parse_xlsx


# ---------------------------------------------------------------------------
# Extrato
# ---------------------------------------------------------------------------

CSV_WITH_PREAMBLE = (
    "Titular;Empresa Teste\n"
    "IBAN;PT50 0000 0000 0000 0000 0000 0\n"
    "\n"
    "Data Mov.;Data Valor;Descrição;Débito;Crédito;Saldo\n"
    "01/10/2026;01/10/2026;RENDA OUTUBRO;850,00;;1.200,00\n"
    "02/10/2026;02/10/2026;TRF CLIENTE;;1.234,56;2.434,56\n"
    ";;Saldo final;;;2.434,56\n"
)


def _xlsx_statement() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["Extrato de conta à ordem"])
    ws.append([])
    ws.append(["Data", "Descritivo", "Montante", "Saldo"])
    ws.append([datetime(2026, 10, 3), "PAG ELECTRICIDADE", -45.30, 100])
    ws.append([datetime(2026, 10, 4), "TRF RECEBIDA", 500, 600])
    ws.append([None, "Total", 454.70, None])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_csv_with_account_lines_above_the_header():
    entries = parse_csv(CSV_WITH_PREAMBLE.encode("utf-8"))

    assert [(e.date, e.type, e.amount) for e in entries] == [
        ("2026-10-01", "debit", 850.00),
        ("2026-10-02", "credit", 1234.56),
    ]
    assert entries[1].balance == 2434.56


def test_excel_statement_reads_typed_cells_and_skips_totals():
    entries = parse_xlsx(_xlsx_statement())

    assert [(e.date, e.description, e.type, e.amount) for e in entries] == [
        ("2026-10-03", "PAG ELECTRICIDADE", "debit", 45.30),
        ("2026-10-04", "TRF RECEBIDA", "credit", 500.00),
    ]


def test_excel_statement_upload_creates_the_movements(tenant):
    response = tenant.post(
        "/api/v1/bank/upload",
        files={"file": ("extrato-cgd.xlsx", _xlsx_statement(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200, response.text
    assert response.json()["total_entries"] == 2


def test_a_broken_excel_file_is_a_clear_error(tenant):
    response = tenant.post(
        "/api/v1/bank/upload",
        files={"file": ("extrato.xlsx", b"isto nao e um excel", "application/octet-stream")},
    )
    assert response.status_code == 400
    assert "Excel" in response.json()["detail"]


# ---------------------------------------------------------------------------
# Exportar
# ---------------------------------------------------------------------------

TABLE = {
    "title": "Contas a Pagar",
    "subtitle": "Em aberto",
    "columns": [
        {"key": "due", "label": "Vencimento", "type": "date"},
        {"key": "desc", "label": "Descrição", "type": "text"},
        {"key": "amount", "label": "Valor", "type": "money"},
    ],
    "rows": [
        {"due": "2026-10-05", "desc": "Renda — escritório", "amount": 850},
        {"due": "2026-10-10", "desc": "Electricidade", "amount": 45.3},
    ],
    "totals": {"desc": "Total", "amount": 895.3},
}


def test_excel_export_keeps_numbers_and_dates_typed(tenant):
    response = tenant.post("/api/v1/reports/export/table", json={**TABLE, "format": "xlsx"})

    assert response.status_code == 200, response.text
    assert "contas-a-pagar" in response.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(response.content)).active
    assert ws["A1"].value == "Contas a Pagar"
    assert "Empresa Teste" in ws["A2"].value
    assert [c.value for c in ws[5]] == ["Vencimento", "Descrição", "Valor"]
    assert ws["A6"].value.date().isoformat() == "2026-10-05"
    assert ws["C6"].value == 850
    assert ws["C8"].value == 895.3  # linha de totais


def test_pdf_export_is_a_pdf_even_with_euro_and_accents(tenant):
    body = {**TABLE, "format": "pdf", "title": "Relatório — Conciliação €"}
    response = tenant.post("/api/v1/reports/export/table", json=body)

    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")


def test_export_needs_a_session(client):
    response = client.post("/api/v1/reports/export/table", json=TABLE)
    assert response.status_code == 401
